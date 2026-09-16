from calendar import monthrange
from datetime import date, time
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import GridRegion, SheetModel, WorkbookModel
from .time_parser import find_time_ranges, parse_time_range


MONTHS = {
    "août": 8,
    "aout": 8,
    "septembre": 9,
    "octobre": 10,
    "novembre": 11,
    "décembre": 12,
    "decembre": 12,
    "janvier": 1,
    "février": 2,
    "fevrier": 2,
}
PAUSE_TEXT = "temps de pause"
DAY_START = time(8, 45)
DAY_END = time(18, 0)
LUNCH_START = time(12, 0)
LUNCH_END = time(14, 0)
MONTH_RE = re.compile(r"^(?P<month>[^ ]+)\s+(?P<year>\d{4})$", re.IGNORECASE)
FULL_DAY_KEYWORDS = ("jour", "journée", "journee", "FÉRIÉ", "FERIÉ", "FERIE", "semaine")


def _month_header(value: object) -> tuple[int, int] | None:
    if not isinstance(value, str):
        return None
    match = MONTH_RE.match(value.strip())
    if not match:
        return None
    month = MONTHS.get(match.group("month").casefold())
    return (month, int(match.group("year"))) if month else None


def _fill_color(cell: object) -> str | None:
    color = cell.fill.fgColor
    if color.type == "rgb" and color.rgb not in (None, "00000000"):
        return color.rgb
    return None


def _merged_range_for(ws: object, coordinate: str) -> object | None:
    return next((item for item in ws.merged_cells.ranges if coordinate in item), None)


def _region_text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _is_structural_value(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _title_from_text(text: str) -> str | None:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if lines and parse_time_range(lines[0]) != (None, None):
        lines = lines[1:]
    return lines[0] if lines else None


def _is_pause(text: str) -> bool:
    return text.casefold().strip() == PAUSE_TEXT


def _mentions_afternoon(text: str) -> bool:
    normalized = text.casefold().replace("’", "'")
    return "après-midi" in normalized or "apres-midi" in normalized


def _mentions_full_day(text: str) -> bool:
    normalized = text.casefold().replace("’", "'")
    keywords = "|".join(keyword.casefold() for keyword in FULL_DAY_KEYWORDS)
    return bool(re.search(rf"\b(?:{keywords})\b", normalized))


def read_workbook(path: str | Path) -> WorkbookModel:
    workbook = load_workbook(path, data_only=False)
    return WorkbookModel(tuple(_read_sheet(ws) for ws in workbook.worksheets))


def _read_sheet(ws: object) -> SheetModel:
    headers = [
        (cell.column, *_month_header(cell.value))
        for cell in ws[1]
        if _month_header(cell.value) is not None
    ]
    regions: list[GridRegion] = []
    region_index = 0
    for header_index, (header_column, month, year) in enumerate(headers):
        next_header = headers[header_index + 1][0] if header_index + 1 < len(headers) else ws.max_column + 1
        day_column = header_column + 1
        content_start = header_column + 3
        for row in range(3, min(ws.max_row, 33) + 1):
            day_value = ws.cell(row, day_column).value
            if not isinstance(day_value, int) or not 1 <= day_value <= monthrange(year, month)[1]:
                continue
            for column in range(content_start, next_header):
                cell = ws.cell(row, column)
                if _is_structural_value(cell.value):
                    continue
                text = _region_text(cell.value)
                if not text:
                    continue
                merged = _merged_range_for(ws, cell.coordinate)
                if merged is not None and cell.coordinate != merged.start_cell.coordinate:
                    continue
                min_row = merged.min_row if merged else row
                max_row = merged.max_row if merged else row
                min_column = merged.min_col if merged else column
                max_column = merged.max_col if merged else column
                region_index += 1
                region_id = f"{ws.title}:R{region_index:04d}"
                regions.append(
                    GridRegion(
                        id=region_id,
                        sheet_name=ws.title,
                        month=month,
                        year=year,
                        day=day_value,
                        min_row=min_row,
                        max_row=max_row,
                        min_column=min_column,
                        max_column=max_column,
                        text=text,
                        source_cells=tuple(
                            f"{get_column_letter(col)}{item_row}"
                            for item_row in range(min_row, max_row + 1)
                            for col in range(min_column, max_column + 1)
                        ),
                        background_color=_fill_color(cell),
                        column_width=sum(
                            (ws.column_dimensions[get_column_letter(col)].width or 8.43)
                            for col in range(min_column, max_column + 1)
                        ),
                    )
                )
    return SheetModel(ws.title, tuple(regions))


def detect_events(model: WorkbookModel) -> tuple:
    from .models import EventCandidate

    events: list[EventCandidate] = []
    for sheet in model.sheets:
        regions_by_day: dict[date, list[GridRegion]] = {}
        for region in sheet.regions:
            region_date = date(region.year, region.month, region.day)
            if region_date.weekday() == 5:
                continue
            regions_by_day.setdefault(region_date, []).append(region)
        for region in sheet.regions:
            region_date = date(region.year, region.month, region.day)
            if region_date.weekday() == 5:
                continue
            if _is_pause(region.text):
                continue
            ranges = find_time_ranges(region.text)
            detected_ranges = ranges or [(None, None, -1, -1)]
            title = _title_from_text(region.text)
            day_regions = regions_by_day[region_date]
            measured_widths = sorted(
                item.column_width for item in day_regions
                if item.column_width is not None and not _is_pause(item.text)
            )
            median_width = measured_widths[len(measured_widths) // 2] if measured_widths else None
            narrow_region = (
                median_width is not None
                and region.column_width is not None
                and region.column_width < median_width * 0.65
            )
            day_ranges = [
                (item, found_range)
                for item in day_regions
                if not _is_pause(item.text)
                for found_range in find_time_ranges(item.text)
            ]
            pauses_before = [
                item for item in day_regions
                if _is_pause(item.text) and item.max_column < region.min_column
            ]
            pause_after_morning = max(pauses_before, key=lambda item: item.max_column, default=None)
            timed_before = [
                found_range for item, found_range in day_ranges
                if item.max_column < region.min_column
            ]
            timed_after = [
                found_range for item, found_range in day_ranges
                if item.min_column > region.max_column
            ]
            afternoon_after_pause = [
                found_range
                for item, found_range in day_ranges
                if pause_after_morning is not None
                and item.min_column > pause_after_morning.max_column
                and found_range[0] >= LUNCH_END
            ]
            afternoon_marker_before = any(
                _mentions_afternoon(item.text) and item.max_column < region.min_column
                for item in day_regions
            )
            for range_index, (start_time, end_time, _, _) in enumerate(detected_ranges, start=1):
                warnings = []
                if _mentions_full_day(region.text) and (start_time is None or end_time is None):
                    start_time = DAY_START
                    end_time = DAY_END
                elif start_time is None:
                    last_morning_end = max(
                        (end for _, end, _, _ in timed_before if end is not None and end <= LUNCH_END),
                        default=None,
                    )
                    morning_finished = bool(timed_before) and all(
                        end is not None and end <= LUNCH_END for _, end, _, _ in timed_before
                    ) and not timed_after
                    has_afternoon_context = bool(pauses_before) or afternoon_marker_before or morning_finished
                    if narrow_region and last_morning_end is not None:
                        start_time = last_morning_end
                        end_time = min(
                            (start for start, _, _, _ in afternoon_after_pause),
                            default=LUNCH_END,
                        )
                    elif has_afternoon_context:
                        start_time = LUNCH_END
                    elif timed_after and not timed_before:
                        start_time = DAY_START
                    else:
                        warnings.append("Impossible de déterminer automatiquement l'heure de début.")
                if end_time is None:
                    if afternoon_after_pause:
                        end_time = min(start for start, _, _, _ in afternoon_after_pause)
                    else:
                        end_time = DAY_END
                confidence = 0.10 + (0.25 if start_time else 0) + (0.25 if end_time else 0) + 0.10
                event_id = region.id.replace(":R", ":E")
                if len(detected_ranges) > 1:
                    event_id = f"{event_id}-{range_index}"
                events.append(
                    EventCandidate(
                        id=event_id,
                        title=title,
                        date=date(region.year, region.month, region.day),
                        start_time=start_time,
                        end_time=end_time,
                        confidence=min(confidence, 1.0),
                        warnings=tuple(warnings),
                        source_region_id=region.id,
                        source_sheet=region.sheet_name,
                        source_range=f"{region.source_cells[0]}:{region.source_cells[-1]}",
                        status="NEEDS_REVIEW" if warnings else "DETECTED",
                    )
                )
    return tuple(events)
